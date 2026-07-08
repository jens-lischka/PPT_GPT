"""infographics.py — true infographics as native, editable PowerPoint shapes.

Everything here is a real freeform (`custGeom`) or auto-shape filled with the OW
palette — no rasters, no SVG at render time. A designer can recolour, resize, or
nudge every piece in PowerPoint. Dispatch via ``draw_infographic(slide, area,
spec)`` where ``spec = {"type": funnel|gauge|venn, ...}``.
"""
from __future__ import annotations

import math
from typing import Any, Dict, Sequence

from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import nsdecls
from pptx.oxml import parse_xml
from runtime import Tokens, PALETTE as _P, FONT_MAJOR as MAJOR, FONT_MINOR as MINOR

EMU = 914400
NAVY = _P["midnightblue"]; SKY = _P["skyblue"]; LBLUE = _P["lblue"]; AMBER = _P["gold"]
GREY = _P["grey"]; PALE = _P["pale"]; CREAM = _P["cream"]; WHITE = _P["white"]
# A navy→light ramp for stacked segments
RAMP = ["000F47", "1B2A63", "3A4E86", "5B73A8", "82BAFF", "CEECFF"]

_SID = [7000]


def _rgb(h):
    return RGBColor.from_string(h.lstrip("#"))


def _freeform(slide, pts, *, fill=None, line=None, line_pt=1.0, alpha=None):
    """Native freeform polygon from inch-space points (closed)."""
    xs = [p[0] for p in pts]; ys = [p[1] for p in pts]
    x0, y0 = min(xs), min(ys)
    w_emu = max(1, int((max(xs) - x0) * EMU)); h_emu = max(1, int((max(ys) - y0) * EMU))

    def _pt(p):
        return f'<a:pt x="{int((p[0]-x0)*EMU)}" y="{int((p[1]-y0)*EMU)}"/>'

    path = ("<a:moveTo>" + _pt(pts[0]) + "</a:moveTo>"
            + "".join("<a:lnTo>" + _pt(p) + "</a:lnTo>" for p in pts[1:])
            + "<a:close/>")
    geom = (f'<a:custGeom><a:avLst/><a:gdLst/><a:ahLst/>'
            f'<a:rect l="0" t="0" r="{w_emu}" b="{h_emu}"/>'
            f'<a:pathLst><a:path w="{w_emu}" h="{h_emu}">{path}</a:path></a:pathLst></a:custGeom>')
    if fill is None:
        fill_xml = "<a:noFill/>"
    elif alpha is not None:
        fill_xml = (f'<a:solidFill><a:srgbClr val="{fill}">'
                    f'<a:alpha val="{int(alpha*1000)}"/></a:srgbClr></a:solidFill>')
    else:
        fill_xml = f'<a:solidFill><a:srgbClr val="{fill}"/></a:solidFill>'
    ln_xml = (f'<a:ln w="{int(line_pt*12700)}"><a:solidFill><a:srgbClr val="{line}"/>'
              f'</a:solidFill></a:ln>') if line else "<a:ln><a:noFill/></a:ln>"
    _SID[0] += 1
    xml = (f'<p:sp {nsdecls("p","a")}><p:nvSpPr>'
           f'<p:cNvPr id="{_SID[0]}" name="infographic-{_SID[0]}"/><p:cNvSpPr/><p:nvPr/>'
           f'</p:nvSpPr><p:spPr><a:xfrm><a:off x="{int(x0*EMU)}" y="{int(y0*EMU)}"/>'
           f'<a:ext cx="{w_emu}" cy="{h_emu}"/></a:xfrm>{geom}{fill_xml}{ln_xml}</p:spPr>'
           f'<p:txBody><a:bodyPr/><a:lstStyle/><a:p/></p:txBody></p:sp>')
    slide.shapes._spTree.append(parse_xml(xml))
    return slide.shapes[-1]


def _embed(shape, lines, *, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE):
    """Write text INTO a shape's own text frame (not a separate overlaid box)."""
    tf = shape.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Pt(4); tf.margin_top = tf.margin_bottom = Pt(2)
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run(); r.text = str(ln.get('text', ''))
        fo = r.font; fo.size = Pt(ln.get('size', 12)); fo.name = ln.get('font', MINOR)
        fo.bold = ln.get('bold', False); fo.color.rgb = _rgb(ln.get('color', NAVY))


def _text(slide, x, y, w, h, lines, *, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = 0            # no fill -> 0 margins
    tf.margin_top = tf.margin_bottom = 0
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run(); r.text = str(ln.get("text", ""))
        f = r.font; f.size = Pt(ln.get("size", 12)); f.name = ln.get("font", MINOR)
        f.bold = ln.get("bold", False); f.color.rgb = _rgb(ln.get("color", NAVY))


# --------------------------------------------------------------------------- #
# Funnel — narrowing trapezoid bands (native freeforms)
# --------------------------------------------------------------------------- #
def funnel(slide, area, stages: Sequence[Dict[str, Any]]) -> None:
    stages = [s for s in stages if isinstance(s, dict)][:6]
    n = len(stages)
    if not n:
        return
    top_w, bot_w = area.w * 0.78, area.w * 0.24
    gap = 0.06
    band_h = (area.h - gap * (n - 1)) / n
    cx = area.x + area.w / 2

    def width_at(frac):
        return top_w + (bot_w - top_w) * frac
    for i, s in enumerate(stages):
        wt = width_at(i / n); wb = width_at((i + 1) / n)
        by = area.y + i * (band_h + gap)
        pts = [(cx - wt / 2, by), (cx + wt / 2, by),
               (cx + wb / 2, by + band_h), (cx - wb / 2, by + band_h)]
        _freeform(slide, pts, fill=RAMP[i % len(RAMP)])
        tc = WHITE if i < 4 else NAVY
        label = s.get("label", ""); value = s.get("value", "")
        _text(slide, cx - wt / 2, by, wt, band_h,
              [{"text": (f"{label}  {value}" if value else label),
                "size": Tokens.TYPE_BODY, "bold": True, "color": tc}])


# --------------------------------------------------------------------------- #
# Gauge — donut ring with a % arc (native shapes + sampled-arc freeform)
# --------------------------------------------------------------------------- #
def _arc_points(cx, cy, r, a0, a1, steps=48):
    return [(cx + r * math.cos(math.radians(a)), cy + r * math.sin(math.radians(a)))
            for a in (a0 + (a1 - a0) * k / steps for k in range(steps + 1))]


def gauge(slide, area, value=0, label="", suffix="%") -> None:
    try:
        pct = max(0.0, min(1.0, float(value) / 100.0))
    except Exception:
        pct = 0.0
    d = min(area.h, area.w) * 0.9
    cx = area.x + area.w / 2; cy = area.y + area.h / 2
    r_out = d / 2; r_in = r_out * 0.62
    # track (full grey ring)
    track = slide.shapes.add_shape(MSO_SHAPE.DONUT,
                                   Inches(cx - r_out), Inches(cy - r_out),
                                   Inches(d), Inches(d))
    track.fill.solid(); track.fill.fore_color.rgb = _rgb(PALE)
    track.line.fill.background(); track.shadow.inherit = False
    try:
        track.adjustments[0] = (r_out - r_in) / r_out
    except Exception:
        pass
    # progress arc (navy annular sector), starting at top, clockwise
    if pct > 0:
        a0, a1 = -90.0, -90.0 + 360.0 * pct
        outer = _arc_points(cx, cy, r_out, a0, a1)
        inner = _arc_points(cx, cy, r_in, a1, a0)
        _freeform(slide, outer + inner, fill=NAVY)
    _text(slide, cx - r_in, cy - r_in, r_in * 2, r_in * 2,
          [{"text": f"{int(round(pct*100))}{suffix}", "size": Tokens.TYPE_KPI, "font": MAJOR, "color": NAVY},
           {"text": label, "size": Tokens.TYPE_BODY, "color": GREY}])


# --------------------------------------------------------------------------- #
# Venn — overlapping translucent circles (native ovals)
# --------------------------------------------------------------------------- #
def venn(slide, area, sets: Sequence[Dict[str, Any]]) -> None:
    sets = [s for s in sets if isinstance(s, dict)][:3]
    n = len(sets)
    if n < 2:
        return
    cols = [NAVY, SKY, AMBER]
    r = min(area.h, area.w) * (0.30 if n == 3 else 0.32)
    cy0 = area.y + area.h / 2
    if n == 2:
        centers = [(area.x + area.w / 2 - r * 0.55, cy0),
                   (area.x + area.w / 2 + r * 0.55, cy0)]
    else:
        centers = [(area.x + area.w / 2, cy0 - r * 0.5),
                   (area.x + area.w / 2 - r * 0.6, cy0 + r * 0.5),
                   (area.x + area.w / 2 + r * 0.6, cy0 + r * 0.5)]
    for (cx, cy), s, col in zip(centers, sets, cols):
        ov = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - r), Inches(cy - r),
                                    Inches(2 * r), Inches(2 * r))
        ov.shadow.inherit = False
        ov.fill.solid(); ov.fill.fore_color.rgb = _rgb(col)
        ov.line.color.rgb = _rgb(NAVY); ov.line.width = Pt(0.75)
        # 45% alpha so overlaps read
        sf = ov.fill.fore_color._xFill.find(
            "{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr")
        if sf is not None:
            sf.append(parse_xml('<a:alpha xmlns:a="http://schemas.openxmlformats.org'
                                '/drawingml/2006/main" val="45000"/>'))
        _embed(ov, [{"text": s.get("label", ""), "size": Tokens.TYPE_BODY, "bold": True,
                     "color": NAVY}], anchor=MSO_ANCHOR.MIDDLE)


def draw_infographic(slide, area, spec: Dict[str, Any]) -> None:
    spec = spec or {}
    t = str(spec.get("type", "")).lower()
    if t == "funnel":
        funnel(slide, area, spec.get("stages", []))
    elif t == "gauge":
        gauge(slide, area, spec.get("value", 0), spec.get("label", ""),
              spec.get("suffix", "%"))
    elif t == "venn":
        venn(slide, area, spec.get("sets", []))
    elif t == "heatmap":
        heatmap(slide, area, spec)


# --------------------------------------------------------------------------- #
# Heatmap — grid of native rectangles, axis labels, stepped legend
# --------------------------------------------------------------------------- #
def _rect(slide, x, y, w, h, *, fill=None, line=None, line_pt=1.0):
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                Inches(x), Inches(y), Inches(w), Inches(h))
    sp.shadow.inherit = False
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = _rgb(fill)
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = _rgb(line); sp.line.width = Pt(line_pt)
    return sp


# Sequential low→high ramp anchors (cream → light blue → sky → navy).
_HEAT_ANCHORS = ["F7F3EE", "CEECFF", "82BAFF", "1B3A7A", "000F47"]


def _lerp_hex(a, b, t):
    return "".join(f"{int(round(int(a[i:i+2],16)+(int(b[i:i+2],16)-int(a[i:i+2],16))*t)):02X}"
                   for i in (0, 2, 4))


def _heat_ramp(n):
    n = max(2, n)
    out = []
    segs = len(_HEAT_ANCHORS) - 1
    for k in range(n):
        f = k / (n - 1) * segs
        i = min(segs - 1, int(f))
        out.append(_lerp_hex(_HEAT_ANCHORS[i], _HEAT_ANCHORS[i + 1], f - i))
    return out


def heatmap(slide, area, spec) -> None:
    """A heatmap of touching native rectangles. spec:
    {values:[[...]], x_labels?, y_labels?, steps?, legend_labels?, description?,
     bg? (cell outline = slide background, default white)}."""
    values = [r for r in (spec.get("values") or []) if isinstance(r, list)]
    if not values:
        return
    rows = len(values); cols = max(len(r) for r in values)
    x_labels = spec.get("x_labels") or []
    y_labels = spec.get("y_labels") or []
    steps = int(spec.get("steps", 5))
    legend_labels = spec.get("legend_labels") or []
    description = spec.get("description", "")
    bg = str(spec.get("bg", "FFFFFF")).lstrip("#")          # cell outline = slide bg
    GAP = 0.25                                              # matrix↔legend spacing rule

    flat = [v for r in values for v in r if isinstance(v, (int, float))]
    if not flat:
        return
    vmin, vmax = min(flat), max(flat)
    ramp = _heat_ramp(steps)

    yl_w = 0.95 if y_labels else 0.0       # room for "N. America" at 12pt
    xl_h = 0.32 if x_labels else 0.0       # one line of 12pt header
    legend_h = 0.42
    mx = area.x + yl_w
    my = area.y + xl_h
    avail_w = area.w - yl_w
    avail_h = area.h - xl_h - GAP - legend_h
    # Rectangle cells that FILL the area (not square) — width and height are
    # independent so the matrix spans the full column instead of bunching into
    # small squares top-left. This also widens columns enough that headers like
    # "Charging" sit on one line instead of wrapping a letter per row.
    cell_w = max(0.30, avail_w / cols)
    cell_h = max(0.22, avail_h / rows)
    mw, mh = cell_w * cols, cell_h * rows

    def bucket(v):
        if vmax == vmin:
            return 0
        return min(steps - 1, int((v - vmin) / (vmax - vmin) * steps))

    # cells (touching, 1pt outline in the slide-background colour)
    for ri, row in enumerate(values):
        for ci in range(cols):
            v = row[ci] if ci < len(row) else None
            fill = ramp[bucket(v)] if isinstance(v, (int, float)) else "FFFFFF"
            _rect(slide, mx + ci * cell_w, my + ri * cell_h, cell_w, cell_h,
                  fill=fill, line=bg, line_pt=1.0)
    # x-axis labels (above) — 12pt, Text 1
    for ci, lab in enumerate(x_labels[:cols]):
        _text(slide, mx + ci * cell_w, area.y, cell_w, xl_h,
              [{"text": str(lab), "size": Tokens.TYPE_BODY, "color": NAVY}],
              align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.BOTTOM)
    # y-axis labels (left, right-aligned to the matrix) — 12pt, Text 1
    for ri, lab in enumerate(y_labels[:rows]):
        _text(slide, area.x, my + ri * cell_h, yl_w - 0.10, cell_h,
              [{"text": str(lab), "size": Tokens.TYPE_BODY, "color": NAVY}],
              align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)

    # legend — left-aligned with the matrix, 0.25" below it. Small chips
    # (sample style) with the label to the right, in 12pt Text 1.
    ly = my + mh + GAP
    sw = 0.16                              # chip size — matches the samples
    cellpitch = max(sw + 0.85, mw / steps)
    for i in range(steps):
        lx = mx + i * cellpitch
        _rect(slide, lx, ly + 0.02, sw, sw, fill=ramp[i], line=bg, line_pt=1.0)
        lab = (legend_labels[i] if i < len(legend_labels)
               else (f"{vmin + (vmax - vmin) * i / steps:.0f}–"
                     f"{vmin + (vmax - vmin) * (i + 1) / steps:.0f}"))
        _text(slide, lx + sw + 0.08, ly - 0.04, cellpitch - sw - 0.14, sw + 0.14,
              [{"text": str(lab), "size": Tokens.TYPE_BODY, "color": NAVY}],
              align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE)
    if description:
        _text(slide, mx, ly + sw + 0.16, mw, 0.3,
              [{"text": str(description), "size": Tokens.TYPE_FOOTNOTE, "color": GREY}],
              align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP)
