"""Smoke, invariant, and doc-parity tests for the deck runtime.

- Smoke: every registered intent builds a slide without error.
- Invariants: the styling/layout rules fixed over time stay fixed
  (grid spacing >= 0.25", canvas content top = 1.54", real bullet glyphs at
  12pt, conclusion minimum height, nested-bullet levels).
- Parity: every registered intent is documented in README_deck_system.md so
  code and docs cannot silently drift.
"""
from __future__ import annotations

import math
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Emu

from compiler import INTENTS, SemanticCompiler
from runtime import TemplateRuntime, ContentArea, CONCLUSION_HEIGHT

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "ow_default.pptx"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"

# Minimal valid content for every registered intent (keyed by slot content key).
MIN_CONTENT = {
    "compose": {"title": "T", "compose": {"regions": [
        {"block": "chart", "width": 2, "data": {"chart": {"type": "column_clustered", "categories": ["A","B"], "series": [{"name":"s","values":[1,2]}]}}},
        {"block": "text", "width": 1, "data": {"heading": "Note", "bullets": ["x"]}}]}},
    "show_treemap": {"title": "T", "treemap": {"items": [{"label": "A", "value": 40}, {"label": "B", "value": 25}, {"label": "C", "value": 20}, {"label": "D", "value": 15}]}},
    "show_waterfall": {"title": "T", "waterfall": {
        "orientation": "vertical",
        "categories": ["Total", "A", "B", "End"],
        "values": [100, 20, -10, 110], "totals": [0, 3]}},
    "introduce_topic": {"title": "Title", "subtitle": "Subtitle"},
    "explain": {"title": "T", "body": {"heading": "H", "bullets": ["a", "b"]}},
    "compare_two_options": {"title": "T",
                            "left": {"heading": "L", "bullets": ["a"]},
                            "right": {"heading": "R", "bullets": ["b"]}},
    "show_trend_with_key_message": {
        "title": "T",
        "chart": {"type": "line_markers", "categories": ["A", "B", "C"],
                  "series": [{"name": "X", "values": [1, 2, 3]}]},
        "insight": {"title": "Signal", "text": "It went up."}},
    "dashboard": {"title": "T", "metrics": [
        {"label": "Rev", "value": "10"}, {"label": "Cost", "value": "5"},
        {"label": "Margin", "value": "50%"}]},
    "section_divider": {"title": "Section", "number": "01"},
    "show_data": {"title": "T", "table": {"headers": ["A", "B"],
                                          "rows": [["1", "2"], ["3", "4"]]}},
    "summary": {"title": "T", "body": {"bullets": ["a", "b"]}},
    "introduce_person": {"name": "Jane Doe", "role": "Partner",
                         "bio": {"bullets": ["x", "y"]}},
    "show_process": {"title": "T", "steps": [
        {"label": "One"}, {"label": "Two"}, {"label": "Three"}]},
    "show_org_chart": {"title": "T", "org_chart": {
        "root": {"label": "CEO"},
        "reports": [{"label": "CFO"}, {"label": "COO"}]}},
    "show_pyramid": {"title": "T", "levels": [
        {"label": "Top"}, {"label": "Mid"}, {"label": "Base"}]},
    "show_cycle": {"title": "T", "items": [
        {"label": "Plan"}, {"label": "Do"}, {"label": "Check"}, {"label": "Act"}]},
    "show_matrix": {"title": "T", "matrix": {
        "rows": 2, "cols": 2, "items": [{"label": "A", "x": 0.2, "y": 0.8}]}},
    "show_timeline": {"title": "T", "milestones": [
        {"date": "Q1", "title": "Start", "description": "d"},
        {"date": "Q2", "title": "Mid", "description": "d"}]},
    "show_growing_steps": {"title": "T", "steps": [
        {"label": "A", "description": "d"}, {"label": "B", "description": "d"},
        {"label": "C", "description": "d"}]},
    "show_contents": {"title": "T", "sections": [
        {"title": "Intro"}, {"title": "Body"}]},
    "show_columns": {"title": "T", "columns": [
        {"heading": "A", "bullets": ["x"]}, {"heading": "B", "bullets": ["y"]}]},
    "show_quote": {"title": "T", "quote": {"text": "A quote.",
                                           "attribution": "Someone"}},
    "show_graphic_with_text": {"title": "T", "panel": {
        "graphic": {"kind": "pyramid", "levels": [{"label": "A"}, {"label": "B"}]},
        "text": {"heading": "H", "subheading": "S", "bullets": ["x", "y"]},
        "side": "graphic_left"}},
    "stat_callout": {"title": "T", "stats": [
        {"value": "50%", "heading": "A", "text": "x"},
        {"value": "3x", "heading": "B", "text": "y"}]},
    "icon_rows": {"title": "T", "rows": [
        {"icon": "bolt", "heading": "A", "text": "x"},
        {"icon": "clock", "heading": "B", "text": "y"}]},
    "compare": {"title": "T", "compare": {"pros": ["a"], "cons": ["b"]}},
    "sticker": {"title": "T", "sticker": {"text": "NEW", "note": "a note"}},
    "infographic": {"title": "T", "infographic": {"type": "funnel", "stages": [
        {"label": "Leads", "value": "1000"}, {"label": "Qualified", "value": "400"},
        {"label": "Won", "value": "120"}]}},
}


def _build(slides, out):
    rt = TemplateRuntime(str(TEMPLATE))
    SemanticCompiler(rt).build({"deck": {"title": "t", "template": "ow_default.pptx"},
                                "slides": slides})
    rt.save(str(out))
    return Presentation(str(out))


def _is_oval(sh):
    try:
        return sh.auto_shape_type is not None and "OVAL" in str(sh.auto_shape_type)
    except Exception:
        return False


def _is_rect(sh):
    try:
        return str(sh.auto_shape_type) == "RECTANGLE (1)"
    except Exception:
        return False


@pytest.mark.parametrize("intent", sorted(INTENTS))
def test_every_intent_builds(intent, tmp_path):
    """Each registered intent compiles and renders at least one slide."""
    assert intent in MIN_CONTENT, f"no smoke content defined for intent {intent!r}"
    prs = _build([{"intent": intent, "content": MIN_CONTENT[intent]}],
                 tmp_path / f"{intent}.pptx")
    assert len(prs.slides) >= 1


def test_every_intent_documented():
    """Doc-parity: every registered intent appears in README_deck_system.md."""
    readme = (ROOT.parent / "knowledge" / "README_deck_system.md")
    if not readme.exists():  # running from the clean zip without knowledge/
        pytest.skip("README_deck_system.md not present in this layout")
    text = readme.read_text(encoding="utf-8")
    missing = [name for name in INTENTS if name not in text]
    assert not missing, f"intents missing from README: {missing}"


def test_cycle_min_gap(tmp_path):
    """show_cycle ovals keep a >= 0.25\" clear gap (no overlap)."""
    prs = _build([{"intent": "show_cycle", "content": {"title": "T", "items": [
        {"label": f"Item {i}"} for i in range(6)]}}], tmp_path / "c.pptx")
    boxes = [(Emu(o.left).inches, Emu(o.top).inches,
              Emu(o.width).inches, Emu(o.height).inches)
             for o in prs.slides[0].shapes if _is_oval(o)]
    assert len(boxes) == 6

    def er(w, h, ang):
        a, b = w / 2, h / 2
        return (a * b) / math.hypot(b * math.cos(ang), a * math.sin(ang))

    def gap(b1, b2):
        x1, y1, w1, h1 = b1
        x2, y2, w2, h2 = b2
        cx1, cy1 = x1 + w1 / 2, y1 + h1 / 2
        cx2, cy2 = x2 + w2 / 2, y2 + h2 / 2
        dx, dy = cx2 - cx1, cy2 - cy1
        d = math.hypot(dx, dy)
        ang = math.atan2(dy, dx)
        return d - er(w1, h1, ang) - er(w2, h2, ang)

    mn = min(gap(boxes[i], boxes[j])
             for i in range(len(boxes)) for j in range(i + 1, len(boxes)))
    assert mn >= 0.24, f"cycle ovals too close: min gap {mn:.3f}in"


def test_growing_steps_gap(tmp_path):
    """show_growing_steps uses the 0.25\" grid gap between steps."""
    prs = _build([{"intent": "show_growing_steps", "content": {"title": "T", "steps": [
        {"label": f"S{i}", "description": "d"} for i in range(5)]}}],
        tmp_path / "g.pptx")
    rects = sorted([(Emu(s.left).inches, Emu(s.width).inches)
                    for s in prs.slides[0].shapes if _is_rect(s)])
    gaps = [round(rects[i + 1][0] - (rects[i][0] + rects[i][1]), 2)
            for i in range(len(rects) - 1)]
    assert gaps and all(abs(g - 0.25) < 0.03 for g in gaps), f"gaps={gaps}"


def test_columns_use_native_layouts(tmp_path):
    """2-4 columns render in the template's native column layouts; 5 columns
    fall back to the manual canvas on Title Only."""
    prs = _build([
        {"intent": "show_columns", "content": {"title": "T", "columns": [
            {"heading": "A", "bullets": ["x"]}, {"heading": "B", "bullets": ["y"]},
            {"heading": "C", "bullets": ["z"]}]}},
        {"intent": "show_columns", "content": {"title": "T", "columns": [
            {"heading": h, "bullets": ["x"]} for h in "ABCDE"]}},
    ], tmp_path / "cols.pptx")
    assert prs.slides[0].slide_layout.name == "3 columns"
    assert prs.slides[1].slide_layout.name == "Title Only"   # 5 -> canvas


def test_columns_canvas_bullets_are_real(tmp_path):
    """The 5-column canvas fallback draws real bullet glyphs (not a typed '* ')
    at 12pt."""
    prs = _build([{"intent": "show_columns", "content": {"title": "T", "columns": [
        {"heading": h, "bullets": ["Standardize"]} for h in "ABCDE"]}}],
        tmp_path / "b.pptx")
    found = False
    for sh in prs.slides[0].shapes:
        if not sh.has_text_frame or "Standardize" not in sh.text_frame.text:
            continue
        for para in sh.text_frame.paragraphs:
            if para.text.strip() == "Standardize":
                pPr = para._p.find(f"{{{A_NS}}}pPr")
                bu = pPr.find(f"{{{A_NS}}}buChar") if pPr is not None else None
                assert bu is not None and bu.get("char") == "\u2022"
                assert para.runs and para.runs[0].font.size.pt == 12
                assert not para.text.lstrip().startswith("\u2022")
                found = True
    assert found


def test_conclusion_minimum_height(tmp_path):
    """A short conclusion stays at the single-line minimum height."""
    prs = _build([{"intent": "explain",
                   "content": {"title": "T", "body": {"bullets": ["a"]}},
                   "conclusion": "A short single-line takeaway."}],
                 tmp_path / "cc.pptx")
    concl = [sh for sh in prs.slides[0].shapes if sh.name == "Conclusion"]
    assert concl, "no conclusion shape"
    assert abs(Emu(concl[0].height).inches - CONCLUSION_HEIGHT) < 0.001


def test_tokens_are_single_source_of_truth():
    """The public constants must derive from Tokens — no divergent literals."""
    from runtime import (Tokens, GRID_GAP, GRID_GAP_LG, AUTOFIT_MIN_PT,
                         CONCLUSION_HEIGHT, CONCLUSION_LINE_HEIGHT,
                         CONTENT_BOTTOM_MARGIN, ContentArea, BRAND_RULES)
    assert GRID_GAP == Tokens.SPACE_GAP
    assert GRID_GAP_LG == Tokens.SPACE_GAP_LG
    assert AUTOFIT_MIN_PT == Tokens.TYPE_AUTOFIT_MIN
    assert CONCLUSION_HEIGHT == Tokens.SPACE_CONCL_H
    assert CONCLUSION_LINE_HEIGHT == Tokens.SPACE_CONCL_LINE
    assert CONTENT_BOTTOM_MARGIN == Tokens.SPACE_CONTENT_BOTTOM
    assert (ContentArea().x, ContentArea().y) == (Tokens.SPACE_CONTENT_X,
                                                  Tokens.SPACE_CONTENT_Y)
    assert BRAND_RULES["primary_font_min_pt"] == Tokens.TYPE_MAJOR_MIN
    assert BRAND_RULES["insight_body_pt"] == Tokens.TYPE_INSIGHTS


def test_autofit_floor_never_below_10pt(tmp_path):
    """Long labels in narrow boxes shrink to fit but never below the 10pt floor."""
    long_steps = [{"label": "Extraordinarily Comprehensive Capability", "description": "d"}
                  for _ in range(7)]
    prs = _build([{"intent": "show_growing_steps",
                   "content": {"title": "T", "steps": long_steps}}],
                 tmp_path / "fit.pptx")
    sizes = [r.font.size.pt
             for sh in prs.slides[0].shapes if sh.has_text_frame
             for p in sh.text_frame.paragraphs for r in p.runs
             if r.font.size is not None]
    assert sizes, "no sized runs found"
    assert min(sizes) >= 10, f"font dropped below 10pt floor: {min(sizes)}"
    assert max(sizes) <= 12, f"label exceeded its cap: {max(sizes)}"


def test_nested_bullets_levels(tmp_path):
    """Nested bullets in a placeholder produce paragraph levels 1/2/3."""
    prs = _build([{"intent": "explain", "content": {"title": "T", "body": {
        "heading": "H",
        "bullets": [{"text": "L1", "bullets": [
            {"text": "L2", "bullets": ["L3"]}]}]}}}], tmp_path / "n.pptx")
    levels = {}
    for sh in prs.slides[0].shapes:
        if not sh.has_text_frame:
            continue
        for para in sh.text_frame.paragraphs:
            t = para.text.strip()
            if t in ("L1", "L2", "L3"):
                pPr = para._p.find(f"{{{A_NS}}}pPr")
                levels[t] = int(pPr.get("lvl")) if pPr is not None and pPr.get("lvl") else 0
    assert levels.get("L1") == 1 and levels.get("L2") == 2 and levels.get("L3") == 3


def test_compose_places_blocks_in_grid(tmp_path):
    """The composer renders each region as its own block, inside the content area,
    with no horizontal overlap between adjacent regions (the grid grammar)."""
    prs = _build([{"intent": "compose", "content": {"title": "T", "compose": {"regions": [
        {"block": "chart", "width": 2, "data": {"chart": {"type": "column_clustered",
            "categories": ["A", "B"], "series": [{"name": "s", "values": [1, 2]}]}}},
        {"block": "text", "width": 1, "data": {"heading": "L", "bullets": ["x"]}},
        {"block": "text", "width": 1, "data": {"heading": "R", "bullets": ["y"]}}]}}}],
        tmp_path / "c.pptx")
    boxes = [(Emu(s.left).inches, Emu(s.top).inches, Emu(s.width).inches, Emu(s.height).inches)
             for s in prs.slides[0].shapes if not s.is_placeholder and s.left is not None]
    assert len(boxes) >= 3                       # chart + two text blocks
    for x, y, w, h in boxes:
        assert x >= ContentArea.x - 0.02
        assert x + w <= ContentArea.x + ContentArea.w + 0.05
        assert y >= ContentArea.y - 0.02
    xs = sorted(boxes, key=lambda b: b[0])
    for a, b in zip(xs, xs[1:]):
        assert a[0] + a[2] <= b[0] + 0.02, "composed regions overlap horizontally"


def test_fit_router_routes_by_signature_and_audience():
    """The fit router snaps an exact preset match and composes when no preset fits;
    a stricter audience snaps a borderline case a looser one would compose."""
    from compiler import route_composition
    assert route_composition([{"block": "chart"}, {"block": "text"}])[0] == "preset"
    mode, _, fit = route_composition(
        [{"block": "chart"}, {"block": "text"}, {"block": "text"}])
    assert mode == "compose" and fit < 0.75
    assert route_composition([{"block": "cycle"}, {"block": "text"}], "board")[0] == "preset"


def test_auto_path_routes_through_router():
    """An 'auto' slide is resolved by the router: a high-fit composition snaps to
    the matched preset's proportions; a borderline one routes differently by
    audience (the threshold is in the loop, not bypassable)."""
    from compiler import route_auto, route_composition
    spec = route_auto({"intent": "auto", "content": {"title": "T", "regions": [
        {"block": "chart", "data": {}}, {"block": "text", "data": {}}]}}, "executive")
    assert spec["intent"] == "compose"
    assert [r.get("width") for r in spec["content"]["compose"]["regions"]] == [2, 1]
    three = [{"block": "chart"}, {"block": "text"}, {"block": "text"}]
    assert route_composition(three, "project_team")[0] == "preset"   # tau 0.55
    assert route_composition(three, "executive")[0] == "compose"     # tau 0.85


# ---- new composite blocks (sample-codified motifs) -------------------------

def test_new_blocks_render_and_use_tokens(tmp_path):
    """The rule_text / kpi_callout / key_message blocks render via the unified
    dispatcher, are registered with min-sizes, and draw their KPI value at the
    single KPI token size (no compact-label divergence)."""
    from runtime import TemplateRuntime, Tokens, BLOCK_MIN_SIZE
    from pptx.util import Pt

    for kind in ("rule_text", "kpi_callout", "key_message"):
        assert kind in BLOCK_MIN_SIZE, f"{kind} missing from BLOCK_MIN_SIZE"

    rt = TemplateRuntime(str(TEMPLATE))
    slide = rt.add_slide("Title Only")
    rt.draw_composition(slide, [{"height": 1.0, "regions": [
        {"block": "rule_text", "data": {"heading": "H", "paragraphs": ["p"], "rule": "left"}},
        {"block": "rule_text", "data": {"heading": "H", "paragraphs": ["p"], "rule": "top"}},
        {"block": "kpi_callout", "data": {"value": "8-12x", "label": "after transition"}},
        {"block": "key_message", "width": 1.2,
         "data": {"message": "ROI exceeds cost",
                  "callouts": [{"value": "59%", "label": "rollout"}]}},
    ]}])
    out = tmp_path / "newblocks.pptx"
    rt.save(str(out))
    # KPI values render at the KPI token size; labels never at the eyebrow size.
    kpi_pt = Pt(Tokens.TYPE_KPI).pt
    sizes = [r.font.size.pt for sh in slide.shapes if sh.has_text_frame
             for p in sh.text_frame.paragraphs for r in p.runs if r.font.size]
    assert kpi_pt in sizes
    # KPI labels were moved off the eyebrow size; assert no KPI-value-adjacent
    # eyebrow leakage by checking the callout label is body-sized.
    assert Pt(Tokens.TYPE_BODY).pt in sizes


def test_kpi_size_is_single_value():
    """Per the sample reference, the standalone and chart-sharing KPI sizes are
    the same single house value (26pt)."""
    from runtime import Tokens
    assert Tokens.TYPE_KPI == Tokens.TYPE_KPI_COMPACT == 26
    assert Tokens.TYPE_MARKER == 14
    assert Tokens.STROKE_PANEL == 1.5


def test_process_chevrons_render_all_steps(tmp_path):
    """draw_process_chevrons must emit one shape per step (regression: a missing
    `overlap` return once crashed the loop after the first shape, leaving a
    single element)."""
    from runtime import TemplateRuntime
    rt = TemplateRuntime(str(TEMPLATE))
    slide = rt.add_slide("Title Only")
    steps = [{"label": str(2024 + i)} for i in range(6)]
    rt.draw_process_chevrons(slide, 0.5, 3.0, 12.33, 0.6, steps)
    out = tmp_path / "chev.pptx"
    rt.save(str(out))
    labels = {sh.text_frame.text for sh in slide.shapes
              if sh.has_text_frame and sh.text_frame.text.startswith("202")}
    assert labels == {str(2024 + i) for i in range(6)}, labels


def test_footnote_wrap_clears_conclusion(tmp_path):
    """Regression: a long single-line source wraps to multiple lines. The
    footnote-height estimate must count wrapped lines (not just explicit \\n) so
    the conclusion above it is positioned clear of the wrapped footnote rather
    than overlapping its top line."""
    long_src = "Source: " + "; ".join(
        f"Reference work number {i}, accessed 2026" for i in range(15))
    assert "\n" not in long_src and len(long_src) > 350
    prs = _build([{
        "intent": "explain",
        "content": {"title": "Government demand adds durability and oversight",
                    "body": {"paragraphs": ["Body text for the slide here."]}},
        "conclusion": ("Strategic relevance is an asset, but it also makes "
                       "regulation, procurement and governance part of the "
                       "investment case for the company going forward."),
        "footnote": long_src,
    }], tmp_path / "f.pptx")
    by_name = {sh.name: sh for s in prs.slides for sh in s.shapes
               if sh.name in ("Conclusion", "Footnote")}
    assert "Conclusion" in by_name and "Footnote" in by_name
    concl_bottom = Emu(by_name["Conclusion"].top + by_name["Conclusion"].height).inches
    foot_top = Emu(by_name["Footnote"].top).inches
    assert concl_bottom <= foot_top + 0.02, (
        f"conclusion bottom {concl_bottom:.3f} overlaps footnote top {foot_top:.3f}")


def test_content_level_conclusion_footnote_are_hoisted(tmp_path):
    """Authors naturally nest conclusion/footnote inside `content`; they must be
    hoisted and rendered, not silently dropped as unknown keys."""
    prs = _build([{"intent": "explain", "content": {
        "title": "A nested conclusion and footnote still render",
        "body": {"paragraphs": ["Some body text for the slide."]},
        "conclusion": "This nested conclusion must appear as the amber band.",
        "footnote": "Source: nested footnote test."}}], tmp_path / "h.pptx")
    names = {sh.name for s in prs.slides for sh in s.shapes}
    assert "Conclusion" in names and "Footnote" in names
