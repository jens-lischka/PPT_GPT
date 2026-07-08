"""Tests for the flexible grid solver and spanning layout."""
from __future__ import annotations

import sys
from pathlib import Path

import pytest
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
TEMPLATE = ROOT / "ow_default.pptx"

from compiler import SemanticCompiler                       # noqa: E402
from runtime import TemplateRuntime, solve_grid, Tokens     # noqa: E402

EMU = 914400.0


def _build(cells, out, cols=None, rows=None):
    grid = {"cells": cells}
    if cols:
        grid["cols"] = cols
    if rows:
        grid["rows"] = rows
    rt = TemplateRuntime(str(TEMPLATE))
    SemanticCompiler(rt).build({"deck": {"title": "t", "template": "ow_default.pptx"},
                                "slides": [{"intent": "compose", "content": {
                                    "title": "Grid layout test slide here",
                                    "compose": {"grid": grid}}}]})
    rt.save(str(out))
    return Presentation(str(out))


@pytest.mark.parametrize("k,expected", [
    (1, (1, 1)), (2, (2, 1)), (3, (3, 1)), (4, (4, 1)),
    (5, (3, 2)), (6, (3, 2)), (7, (4, 2)), (8, (4, 2)), (9, (3, 3)),
])
def test_solve_grid_minimises_rows(k, expected):
    assert solve_grid(k) == expected


def test_grid_spanning_geometry(tmp_path):
    """A 3x2 grid: a 2-col span is 2*col_w + gutter; a 2-row span is full height."""
    prs = _build([
        {"block": "text", "col": 0, "row": 0, "colspan": 2,
         "data": {"paragraphs": ["span two columns"]}},
        {"block": "text", "col": 2, "row": 0, "rowspan": 2,
         "data": {"paragraphs": ["span two rows"]}},
        {"block": "text", "col": 0, "row": 1, "data": {"paragraphs": ["a"]}},
        {"block": "text", "col": 1, "row": 1, "data": {"paragraphs": ["b"]}},
    ], tmp_path / "g.pptx", cols=3, rows=2)
    boxes = [(s.left / EMU, s.top / EMU, s.width / EMU, s.height / EMU)
             for s in prs.slides[0].shapes if not s.is_placeholder and s.width / EMU > 1]
    col_w = (12.333 - 2 * 0.5) / 3                       # 3.778
    span2 = 2 * col_w + 0.5                              # 8.056
    full_h = 2 * ((5.060 - 0.5) / 2) + 0.5               # 5.060
    assert any(abs(w - span2) < 0.05 for _, _, w, _ in boxes), "no 2-col span found"
    assert any(abs(h - full_h) < 0.05 for _, _, _, h in boxes), "no 2-row span found"


def test_grid_narrow_gutter(tmp_path):
    """A 6-column grid drops to the 0.25\" gutter (0.5\" would breach min width)."""
    cells = [{"block": "text", "data": {"paragraphs": [str(i)]}} for i in range(6)]
    prs = _build(cells, tmp_path / "n.pptx", cols=6, rows=1)
    xs = sorted(s.left / EMU for s in prs.slides[0].shapes
                if not s.is_placeholder and s.width / EMU > 1)
    ws = [s.width / EMU for s in prs.slides[0].shapes
          if not s.is_placeholder and s.width / EMU > 1]
    # gutter between adjacent columns ~ 0.25"
    gutter = xs[1] - (xs[0] + ws[0])
    assert abs(gutter - 0.25) < 0.05, f"expected 0.25\" gutter, got {gutter:.3f}"
    assert all(w >= Tokens.GRID_MIN_COL_W - 0.05 for w in ws), "columns below min width"


import gate as _gate                                       # noqa: E402


def test_gate_grid_valid_passes():
    deck = {"slides": [{"content": {"compose": {"grid": {"cols": 3, "rows": 2, "cells": [
        {"block": "text", "col": 0, "row": 0, "colspan": 2},
        {"block": "text", "col": 2, "row": 0, "rowspan": 2},
        {"block": "text", "col": 0, "row": 1},
        {"block": "text", "col": 1, "row": 1},
    ]}}}}]}
    assert _gate._check_grid(deck["slides"])["passed"]


def test_gate_grid_overlap_fails():
    deck = {"slides": [{"content": {"compose": {"grid": {"cols": 2, "rows": 2, "cells": [
        {"block": "text", "col": 0, "row": 0, "colspan": 2},
        {"block": "text", "col": 1, "row": 0},          # overlaps the span
    ]}}}}]}
    res = _gate._check_grid(deck["slides"])
    assert not res["passed"] and any("overlap" in f for f in res["fail_items"])


def test_gate_grid_out_of_bounds_fails():
    deck = {"slides": [{"content": {"compose": {"grid": {"cols": 2, "rows": 1, "cells": [
        {"block": "text", "col": 1, "row": 0, "colspan": 2},   # spills past col 2
    ]}}}}]}
    res = _gate._check_grid(deck["slides"])
    assert not res["passed"] and any("declared" in f for f in res["fail_items"])
