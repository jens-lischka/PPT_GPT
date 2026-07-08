"""Tests for the design-system tooling: component inventory, prose merge,
deck lint, contrast check, and token diff."""
import sys, os
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src"))

import runtime as rt
import design_components as dc
import design_lint as dl


# --- token resolver --------------------------------------------------------
def test_resolve_palette_ref():
    assert rt.resolve_color_ref("{palette.midnightblue}") == "000F47"

def test_resolve_chained_color_role():
    # subtle_highlight -> accent3 -> CEECFF
    assert rt.resolve_color_ref("{BRAND_COLORS.subtle_highlight}") == "CEECFF"

def test_resolve_size_token():
    assert rt.resolve_token_ref("{Tokens.TYPE_KPI}") == rt.Tokens.TYPE_KPI

def test_broken_ref_raises():
    import pytest
    with pytest.raises(Exception):
        rt.resolve_token_ref("{palette.does_not_exist}")


# --- component inventory ---------------------------------------------------
def test_every_component_token_ref_resolves():
    """No component may bind a token that doesn't resolve — this is the
    broken-ref guard at the source."""
    for name, spec in dc.COMPONENTS.items():
        for role, ref in spec["tokens"].items():
            rt.resolve_token_ref(ref)   # raises on broken ref

def test_components_appear_in_generated_doc():
    md = rt.design_system_markdown()
    assert "## Components — block inventory" in md
    for name in dc.component_names():
        assert f"### `{name}`" in md


# --- prose merge -----------------------------------------------------------
def test_prose_is_merged_into_doc():
    md = rt.design_system_markdown()
    # a distinctive phrase from design_prose SECTION_PROSE["colors"]
    assert "addressed by *role*" in md

def test_prose_keys_are_known_sections():
    """Guard against prose rotting: every SECTION_PROSE key should be one the
    generator actually emits."""
    import design_prose as dp
    known = {"overview", "type", "space", "para", "inset", "stroke",
             "chart", "bullets", "align", "colors", "rules"}
    assert set(dp.SECTION_PROSE).issubset(known)

def test_component_prose_keys_match_components():
    import design_prose as dp
    for k in dp.COMPONENT_PROSE:
        assert k in dc.COMPONENTS, f"prose for unknown component {k!r}"


# --- contrast --------------------------------------------------------------
def test_contrast_report_passes():
    rep = dl.contrast_report()
    assert rep["passed"], rep["findings"]

def test_contrast_ratio_known_value():
    # navy on white is a high-contrast pair
    assert rt.contrast_ratio("000F47", "FFFFFF") > 15.0


# --- deck lint -------------------------------------------------------------
def test_lint_clean_deck_passes():
    spec = {"slides": [{"intent": "summary", "content": {"title": "A real insight here"}}]}
    rep = dl.lint_deck(spec)
    assert rep["passed"]

def test_lint_flags_unknown_intent():
    spec = {"slides": [{"intent": "totally_made_up", "content": {}}]}
    rep = dl.lint_deck(spec)
    assert not rep["passed"]
    assert any(f["severity"] == "error" for f in rep["findings"])

def test_lint_flags_off_palette_color():
    spec = {"slides": [{"intent": "summary",
                        "content": {"title": "x", "fill": "#FF00FF"}}]}
    rep = dl.lint_deck(spec)
    assert any("not an OW palette" in f["message"] for f in rep["findings"])

def test_lint_accepts_palette_role():
    spec = {"slides": [{"intent": "summary",
                        "content": {"title": "x", "fill": "subtle_highlight"}}]}
    rep = dl.lint_deck(spec)
    assert not any(f["path"].endswith(".fill") and f["severity"] == "warning"
                   for f in rep["findings"])


# --- diff ------------------------------------------------------------------
def test_diff_detects_removed_token_as_regression():
    a = "| `TYPE_FOO` | 12 |\n| `TYPE_BAR` | 14 |"
    b = "| `TYPE_BAR` | 14 |"
    rep = dl.diff_tokens(a, b)
    assert rep["regression"]
    assert "TYPE_FOO" in rep["removed"]

def test_diff_identical_is_clean():
    a = "| `TYPE_FOO` | 12 |"
    rep = dl.diff_tokens(a, a)
    assert not rep["regression"]
    assert not rep["added"] and not rep["removed"] and not rep["modified"]

def test_diff_reports_modification():
    a = "| `SPACE_GAP` | 0.25\" |"
    b = "| `SPACE_GAP` | 0.30\" |"
    rep = dl.diff_tokens(a, b)
    assert not rep["regression"]   # modified, not removed
    assert rep["modified"][0]["token"] == "SPACE_GAP"


# --- colour section rework (v2.4.1) ----------------------------------------
def test_color_roles_resolve_to_literal_hex():
    """The semantic-role table must show literal hex, not raw accentN."""
    md = rt.design_system_markdown()
    # subtle_highlight is stored as accent3; the doc must show #CEECFF, with
    # accent3 only as provenance in the Via column.
    assert "#CEECFF" in md
    assert "| `subtle_highlight` | `accent3` |" not in md  # no raw-indirection cell

def test_raw_palette_table_present_at_top():
    md = rt.design_system_markdown()
    assert "**OW palette (literal)**" in md
    # colour section precedes the type scale (colour leads the doc)
    assert md.index("## Colours") < md.index("## Type scale")
    for name in rt.PALETTE:
        assert f"`{name}`" in md

def test_color_usage_keys_match_brand_colors():
    import design_prose as dp
    for role in dp.COLOR_USAGE:
        assert role in rt.BRAND_COLORS, f"usage for unknown role {role!r}"


# --- closing-density + bullet-cap (deck-feedback fixes) ---------------------
def test_density_flags_parallel_paragraphs_no_visual():
    import sys, os
    sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src"))
    from gate import _check_density
    slide = {"intent": "explain", "content": {"title": "Modular reports",
             "body": {"paragraphs": ["Use CEO when…", "Use CFO when…", "Use Cities when…"]}}}
    r = _check_density([slide])
    assert r["fail_items"] and "show_columns" in r["fail_items"][0]

def test_density_quiet_when_visual_present():
    import sys, os
    sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src"))
    from gate import _check_density
    slide = {"intent": "explain", "content": {"title": "x",
             "body": {"paragraphs": ["a", "b", "c"]}, "chart": {"type": "line"}}}
    assert not _check_density([slide])["fail_items"]

def test_consulting_bullet_cap_is_30():
    import sys, os
    sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src"))
    import voice_profiles as vp
    assert vp.VOICE_PROFILES["consulting"]["body"]["bullets_max_per_unit"] == 30


# --- chart caption as text shape + data labels (deck-feedback fixes) --------
def _build_dashboard_with_captioned_chart(tmp_path):
    import sys, os
    sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src"))
    import runtime
    tmpl = os.path.join(os.path.dirname(__file__), "..", "ow_default.pptx")
    rt = runtime.TemplateRuntime(tmpl)
    spec = {"slides": [{"intent": "dashboard", "content": {"title": "Dash",
        "metrics": [
            {"heading": "Revenue", "subheading": "US$bn",
             "chart": {"type": "column_clustered", "categories": ["A", "B", "C"],
                       "series": [{"name": "R", "values": [1, 2, 3]}]}},
            {"label": "Market", "value": "€420B"}]}}]}
    out = str(tmp_path / "dash.pptx")
    rt.generate(spec, out) if hasattr(rt, "generate") else None
    return rt, spec, out

def test_chart_caption_is_text_shape_not_embedded_title(tmp_path):
    import os, json, subprocess
    src = os.path.join(os.path.dirname(__file__), "..", "src")
    tmpl = os.path.join(os.path.dirname(__file__), "..", "ow_default.pptx")
    spec = {"slides": [{"intent": "dashboard", "content": {"title": "Dash",
        "metrics": [{"heading": "Revenue", "subheading": "US$bn",
             "chart": {"type": "column_clustered", "categories": ["A", "B", "C"],
                       "series": [{"name": "R", "values": [1, 2, 3]}]}}]}}]}
    deck = tmp_path / "d.json"; deck.write_text(json.dumps(spec))
    out = tmp_path / "d.pptx"
    subprocess.run(["python", "generate_deck.py", "--template", tmpl,
                    "--deck", str(deck), "--output", str(out)],
                   cwd=src, capture_output=True)
    from pptx import Presentation
    s = Presentation(str(out)).slides[0]
    captions = [sh.text_frame.text for sh in s.shapes
                if sh.has_text_frame and "Revenue" in sh.text_frame.text]
    charts = [sh.chart for sh in s.shapes if sh.has_chart]
    assert captions and "US$bn" in captions[0]          # caption is a text shape
    assert charts and charts[0].has_title is False        # not embedded in chart
    assert charts[0].plots[0].has_data_labels is True     # values visible


# --- heatmap / timeline / overflow-inset (round-3 deck feedback) ------------
def test_heatmap_labels_are_12pt_text1():
    import sys, os
    sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src"))
    import inspect, infographics
    src = inspect.getsource(infographics.heatmap)
    assert "Tokens.TYPE_BODY" in src and "TYPE_FOOTNOTE" not in src.split("if description")[0]
    assert "cell_w" in src and "cell_h" in src          # rectangle cells, not square
    assert 'sw = 0.16' in src                            # sample-style legend chip

def test_timeline_font_is_12pt():
    import sys, os
    sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src"))
    import inspect, runtime
    src = inspect.getsource(runtime.TemplateRuntime.draw_timeline)
    assert "Tokens.TYPE_BODY" in src                     # 12pt house body

def test_overflow_report_accounts_for_insets():
    import sys, os
    sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src"))
    import inspect, textmetrics
    src = inspect.getsource(textmetrics.overflow_report)
    assert "margin_left" in src and "w_eff" in src       # insets subtracted


# --- fit-aware card grid (layout chosen by fit, not count) ------------------
def test_fit_aware_grid_picks_taller_when_constrained():
    import sys, os
    sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src"))
    import runtime
    rt = runtime.TemplateRuntime(os.path.join(os.path.dirname(__file__), "..", "ow_default.pptx"))
    full = rt.content_area()
    class A: pass
    narrow = A(); narrow.x = full.x + 0.6 * full.w; narrow.y = full.y
    narrow.w = 0.4 * full.w; narrow.h = full.h
    long_card = {"value": "$1.2B", "heading": "Investment required",
                 "body": {"paragraphs": ["Product localisation, channels, working "
                          "capital and launch marketing across five regions."]}}
    short_card = {"label": "ROI", "value": "16%", "delta": "base"}
    # long cards in a narrow area must not stay a 1x4 strip
    assert rt._pick_card_grid(4, [long_card] * 4, narrow, 0.5) != (1, 4)
    # short cards in the same narrow area can stay a single row
    assert rt._pick_card_grid(4, [short_card] * 4, narrow, 0.5) == (1, 4)
    # no content/area → count-based default preserved
    assert rt._pick_card_grid(4) == (1, 4)


# --- curated icon set + alias resolution (3000-icon upload, curated to 500) --
def test_curated_icon_set_and_aliases():
    import sys, os, glob
    sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src"))
    import icons
    n = len(glob.glob(os.path.join(os.path.dirname(icons.__file__), "assets", "icons", "*.svg")))
    assert n >= 450                                   # curated ~500 set installed
    # the names silently skipped in the field log must now resolve
    for name in ("sparkles", "heart", "device-phone-mobile", "user-group", "chart-bar",
                 "light-bulb", "magnifying-glass", "shield-check", "currency-dollar"):
        assert icons.resolve(name), f"{name} should resolve via alias/fuzzy"


# --- chevron height is a hard 1.0" everywhere ------------------------------
def test_chevron_height_is_hard_one_inch(tmp_path):
    import os, json, subprocess
    src = os.path.join(os.path.dirname(__file__), "..", "src")
    tmpl = os.path.join(os.path.dirname(__file__), "..", "ow_default.pptx")
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE
    cases = {
        "plain": {"slides": [{"intent": "show_process", "content": {"title": "P",
            "steps": [{"label": "A"}, {"label": "B"}, {"label": "C"}]}}]},
        "bullets": {"slides": [{"intent": "show_process", "content": {"title": "P",
            "steps": [{"label": "A", "heading": "1", "bullets": ["x"]},
                      {"label": "B", "heading": "2", "bullets": ["y"]}]}}]},
    }
    for name, spec in cases.items():
        deck = tmp_path / f"{name}.json"; deck.write_text(json.dumps(spec))
        out = tmp_path / f"{name}.pptx"
        subprocess.run(["python", "generate_deck.py", "--template", tmpl,
                        "--deck", str(deck), "--output", str(out)],
                       cwd=src, capture_output=True)
        hs = []
        for s in Presentation(str(out)).slides:
            for sh in s.shapes:
                try:
                    if sh.auto_shape_type in (MSO_SHAPE.CHEVRON, MSO_SHAPE.PENTAGON):
                        hs.append(round(sh.height / 914400, 3))
                except Exception:
                    pass
        assert hs, f"{name}: no chevrons rendered"
        assert all(h == 1.0 for h in hs), f"{name}: heights {hs} != 1.0"


# --- cover layout depends on whether a cover image is present ---------------
def test_cover_layout_follows_image_presence(tmp_path):
    import os, json, subprocess
    from pptx import Presentation
    src = os.path.join(os.path.dirname(__file__), "..", "src")
    tmpl = os.path.join(os.path.dirname(__file__), "..", "ow_default.pptx")
    def cover_layout(content):
        spec = {"slides": [{"intent": "introduce_topic", "content": content}]}
        deck = tmp_path / "c.json"; deck.write_text(json.dumps(spec))
        out = tmp_path / "c.pptx"
        subprocess.run(["python", "generate_deck.py", "--template", tmpl,
                        "--deck", str(deck), "--output", str(out)],
                       cwd=src, capture_output=True)
        return Presentation(str(out)).slides[0].slide_layout.name
    assert cover_layout({"title": "T", "subtitle": "S"}) == "Title Slide"
    assert cover_layout({"title": "T", "subtitle": "S",
                         "image": {"description": "skyline"}}) == "Title Slide with Picture"


# --- charts: negative bars stay filled, x-axis labels pinned to bottom -------
def test_chart_negative_fill_and_axis_labels(tmp_path):
    import os, json, subprocess, zipfile, re
    from lxml import etree
    src = os.path.join(os.path.dirname(__file__), "..", "src")
    tmpl = os.path.join(os.path.dirname(__file__), "..", "ow_default.pptx")
    spec = {"slides": [{"intent": "show_trend_with_key_message", "content": {
        "title": "Neg", "insight": {"heading": "H", "body": "B"},
        "chart": {"type": "column", "categories": ["A", "B", "C"],
                  "series": [{"name": "S1", "values": [-1, -2, -3]},
                             {"name": "S2", "values": [-2, -1, -4]}]}}}]}
    deck = tmp_path / "n.json"; deck.write_text(json.dumps(spec))
    out = tmp_path / "n.pptx"
    subprocess.run(["python", "generate_deck.py", "--template", tmpl,
                    "--deck", str(deck), "--output", str(out)],
                   cwd=src, capture_output=True)
    z = zipfile.ZipFile(str(out))
    cn = [n for n in z.namelist() if re.search(r"charts/chart\d+\.xml$", n)][0]
    ns = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart"}
    root = etree.fromstring(z.read(cn))
    # every bar/column series forces invertIfNegative=0 (solid negative bars)
    inv = [e.get("val") for e in root.iter(f"{{{ns['c']}}}invertIfNegative")]
    assert inv and all(v == "0" for v in inv), f"invertIfNegative={inv}"
    # category axis labels pinned low (bottom), schema-ordered before spPr
    cat = root.find(".//c:catAx", ns)
    tlp = cat.find("c:tickLblPos", ns)
    assert tlp is not None and tlp.get("val") == "low"
    kids = [etree.QName(c).localname for c in cat]
    assert "spPr" not in kids or kids.index("tickLblPos") < kids.index("spPr")


# --- generated decks are authored to the GPT, not the template's creator -----
def test_document_author_is_stamped(tmp_path):
    import os, json, subprocess
    from pptx import Presentation
    src = os.path.join(os.path.dirname(__file__), "..", "src")
    tmpl = os.path.join(os.path.dirname(__file__), "..", "ow_default.pptx")
    spec = {"slides": [{"intent": "introduce_topic",
                        "content": {"title": "T", "subtitle": "S"}}]}
    deck = tmp_path / "a.json"; deck.write_text(json.dumps(spec))
    out = tmp_path / "a.pptx"
    subprocess.run(["python", "generate_deck.py", "--template", tmpl,
                    "--deck", str(deck), "--output", str(out)],
                   cwd=src, capture_output=True)
    cp = Presentation(str(out)).core_properties
    assert cp.author == "CS PresentationGPT"
    assert cp.last_modified_by == "CS PresentationGPT"


# --- tables paginate greedily: max rows per page, remainder continued --------
def test_table_paginates_greedily(tmp_path):
    import os, json, subprocess
    from pptx import Presentation
    src = os.path.join(os.path.dirname(__file__), "..", "src")
    tmpl = os.path.join(os.path.dirname(__file__), "..", "ow_default.pptx")
    rows = [[f"Item {i}", f"{i*10}", f"{i}%"] for i in range(1, 41)]
    spec = {"slides": [{"intent": "show_data", "content": {"title": "Big table",
            "table": {"headers": ["Name", "Value", "Share"], "rows": rows}}}]}
    deck = tmp_path / "t.json"; deck.write_text(json.dumps(spec))
    out = tmp_path / "t.pptx"
    subprocess.run(["python", "generate_deck.py", "--template", tmpl,
                    "--deck", str(deck), "--output", str(out)], cwd=src, capture_output=True)
    p = Presentation(str(out))
    per_slide = [len(sh.table.rows) - 1 for s in p.slides for sh in s.shapes if sh.has_table]
    assert per_slide, "no table rendered"
    assert sum(per_slide) == 40                       # all rows preserved, none cut
    assert per_slide[0] >= 12                          # first page is filled, not ~7
    assert len(per_slide) <= 3                          # 40 rows in <=3 pages (greedy)
    # every page repeats the header and all but the last are filled near-max
    assert all(per_slide[i] >= per_slide[-1] for i in range(len(per_slide) - 1))
    # every table slide carries the header row
    for s in p.slides:
        for sh in s.shapes:
            if sh.has_table:
                assert sh.table.cell(0, 0).text == "Name"


def test_conclusion_height_adjustment_off():
    import sys, os
    sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src"))
    import runtime
    assert runtime.ADJUST_CONCLUSION_HEIGHT is False
    short = runtime.TemplateRuntime.estimate_conclusion_height("Short.")
    long = runtime.TemplateRuntime.estimate_conclusion_height("A " + "very " * 40 + "long one.")
    assert short == long == runtime.CONCLUSION_HEIGHT   # fixed height (no growth)
